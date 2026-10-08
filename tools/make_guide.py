"""Generate the Chinese v1.1 operation manual using python-docx."""
from __future__ import annotations
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NAVY = RGBColor(24, 41, 66)
PURPLE = RGBColor(77, 69, 157)
GRAY = RGBColor(91, 101, 118)
SILVER = RGBColor(239, 242, 248)
WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_font(style, name, size, color=None, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = color
    rp = style.element.get_or_add_rPr()
    rfonts = rp.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rp.insert(0, rfonts)
    for key in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rfonts.set(qn(key), name)


def table(doc, headers, rows, widths=None):
    grid = doc.add_table(rows=1, cols=len(headers))
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.style = 'Table Grid'
    for idx, heading in enumerate(headers):
        cell = grid.rows[0].cells[idx]
        cell.text = heading
        set_cell_shading(cell, '263956')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = WHITE
    for i, data in enumerate(rows):
        row = grid.add_row()
        for j, value in enumerate(data):
            cell = row.cells[j]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2 == 0:
                set_cell_shading(cell, 'F3F5FA')
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs:
                    run.font.size = Pt(9)
    if widths:
        for row in grid.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph('')
    return grid


def note(doc, title, body):
    p = doc.add_paragraph()
    p.style = 'Note'
    r = p.add_run(f'{title}  ')
    r.bold = True
    r.font.color.rgb = PURPLE
    p.add_run(body)


def code(doc, text):
    p = doc.add_paragraph(style='Code')
    p.add_run(text)


def steps(doc, actions):
    for num, action in enumerate(actions, 1):
        p = doc.add_paragraph(f'{num}. {action}')
        p.paragraph_format.left_indent = Cm(.3)


def generate(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(2.15)
    sec.right_margin = Cm(2.15)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(.85)

    normal = doc.styles['Normal']
    set_font(normal, 'Noto Sans CJK SC', 10, NAVY)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.2
    for name, size, spacing in [('Title', 26, 16), ('Heading 1', 16, 12), ('Heading 2', 12.5, 8)]:
        st = doc.styles[name]
        set_font(st, 'Noto Sans CJK SC', size, NAVY, True)
        st.paragraph_format.space_before = Pt(spacing)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    code_style = doc.styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
    set_font(code_style, 'DejaVu Sans Mono', 9, PURPLE)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(9)
    code_style.paragraph_format.left_indent = Cm(.4)
    note_style = doc.styles.add_style('Note', WD_STYLE_TYPE.PARAGRAPH)
    set_font(note_style, 'Noto Sans CJK SC', 9.5, GRAY)
    note_style.paragraph_format.left_indent = Cm(.2)
    note_style.paragraph_format.space_after = Pt(11)

    header = sec.header.paragraphs[0]
    header.text = 'DEV ASSISTANT  /  WINDOWS DEVELOPER WORKSPACE'
    header.style = normal
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = GRAY
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = foot.add_run('Dev Assistant v1.1  ·  中文操作手册  ·  2026-10-08')
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY

    title = doc.add_paragraph(style='Title')
    title.add_run('Dev Assistant\nWindows 操作手册')
    subtitle = doc.add_paragraph('v1.1.0  |  Git 实操学习 · 可输入终端 · 开发工具管理')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in subtitle.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = PURPLE
    doc.add_paragraph('适用系统：Windows 10/11 64 位。目标用户：希望通过图形界面学习 Git，同时管理本机开发工具的开发者。')
    note(doc, '版本提示', '本手册对应 v1.1 在线安装引导器。安装包并不预装 Python / PySide6；首次安装需要网络。构建环境未能进行 Windows 真机 GUI 验收。')
    doc.add_heading('一、快速开始', 1)
    steps(doc, [
        '运行 DevAssistant-Setup-v1.1.exe，阅读安装提示。',
        '安装器检测 Python 3.12 x64；缺失时从官方来源下载安装，并在线安装 PySide6。',
        '安装完成，从桌面或开始菜单启动 Dev Assistant。',
        '在 Projects 页选择“添加项目”，选中本地项目文件夹。',
        '进入 Git 页操作，底部 Terminal 会显示真实命令、输出及退出状态。',
    ])
    table(doc, ['模块', '主要用途', '本版说明'], [
        ('Git', '本地仓库、版本、分支、暂存、远程', '主要学习入口'),
        ('Terminal', '输入命令、查看所有 GUI 执行细节', '单条命令会话，不是完整 PTY'),
        ('Learning', 'Git 命令解释、用法、示例', '离线词典'),
        ('Environment / LFS', '环境检查、依赖导出恢复、大文件追踪', '需本机工具'),
        ('More Tools', 'GitHub / SSH / Docker / 健康检查 / 插件', '按需配置环境'),
    ], [3.6, 6.3, 6.4])

    doc.add_page_break()
    doc.add_heading('二、安装、升级与卸载', 1)
    doc.add_heading('2.1 系统要求与安装', 2)
    table(doc, ['内容', '说明'], [
        ('操作系统', 'Windows 10/11 x64，当前用户安装，无需管理员权限'),
        ('网络', '首次安装/升级需联网安装 Python 及 pip 依赖'),
        ('Git', '建议预装 Git for Windows；需让 git.exe 在 PATH 可找到'),
        ('Docker / Node / Java / LFS', '分别按需要安装，不强制随程序安装'),
        ('默认程序目录', r'%LOCALAPPDATA%\Programs\Dev Assistant'),
        ('配置/历史/日志目录', r'%LOCALAPPDATA%\DevAssistantLab\DevAssistant'),
    ], [5, 11])
    note(doc, '升级提示', '新安装器会在替换旧程序前创建备份；SQLite 存储于用户数据目录，不会在升级时直接覆盖。')
    doc.add_heading('2.2 卸载和故障日志', 2)
    doc.add_paragraph('Windows 设置 → 应用 → 已安装的应用 → Dev Assistant → 卸载。若安装失败，打开以下日志查看错误：')
    code(doc, r'%LOCALAPPDATA%\DevAssistantLab\DevAssistant\install.log')
    doc.add_paragraph('如提示 PySide6 下载失败，检查网络/代理；如提示 Git 不存在，安装 Git for Windows 并重启软件。未签名安装程序可能被 Windows SmartScreen 提示，请先核对下载来源。')

    doc.add_heading('三、Git 图形界面与终端同步学习', 1)
    doc.add_paragraph('原则：点击 Git 页按钮不会悄悄执行“隐藏命令”。底部 Terminal 显示来源标签 [UI · Git]、执行目录、实际 Git 命令、命令输出、Exit Code。')
    code(doc, '[14:02:11] [UI · Git] D:\\Projects\\Demo> git add -A\n'
              '  ⓘ 将文件变更加入暂存区。\n'
              '  ↳ exit 0 · 0.03s\n\n'
              '[14:02:19] [UI · Git] D:\\Projects\\Demo> git commit -m "Fix bug"\n'
              '[main 4d3f42a] Fix bug\n'
              '  ↳ exit 0 · 0.08s')
    doc.add_heading('3.1 从零完成一次本地提交', 2)
    steps(doc, [
        'Projects → 添加并选择一个包含源代码的本地目录。',
        '如果尚未是 Git 仓库，在 Git 页点击“初始化 Git”（命令：git init）。',
        '修改或创建一个文件，点击“刷新状态”（命令：git status --porcelain=v1 -z 等）。',
        '在变更列表选中文件，点击 Stage Selected（命令：git add -- <file>）；也可以 Stage All（git add -A）。',
        '输入清晰的提交信息（例如“修复终端输入”），点击 Commit（git commit -m "..."）。',
        '点击“提交历史”，检查 git log --oneline --decorate --graph -n 50。',
        '在 Terminal 亲自输入 git status 或 git log --oneline，加深记忆。',
    ])
    note(doc, '重点区别', 'git add 只暂存当前内容快照；git commit 将已暂存内容记录为提交。仅点击 Commit 不会自动 Stage All。')
    doc.add_heading('3.2 常用按钮与实际命令对照', 2)
    table(doc, ['GUI 操作', 'Git 命令', '作用'], [
        ('刷新状态', 'git status --porcelain=v1 -z', '机器可解析状态；会显示在 Terminal'),
        ('Stage All', 'git add -A', '暂存所有改动'),
        ('Stage Selected', 'git add -- <file>', '只暂存选中文件'),
        ('Unstage Selected', 'git restore --staged -- <file>', '取消暂存，保留工作区修改'),
        ('Commit', 'git commit -m "message"', '提交暂存内容'),
        ('Diff', 'git diff --no-ext-diff [--cached] -- <file>', '查看工作区/暂存区差异'),
        ('提交历史', 'git log --oneline --decorate --graph -n 50', '查看本地提交'),
        ('Stash / Pop', 'git stash push / git stash pop', '临时保存/恢复变更'),
        ('创建 / 切换分支', 'git branch <name> / git checkout <name>', '创建/切换分支'),
        ('Remote', 'git remote -v / add / remove', '查看或管理远程仓库'),
        ('Git 用户信息', 'git config user.name / user.email', '读取当前仓库的用户身份'),
        ('Pull / Push', 'git pull / git push', '与配置好的远程仓库同步'),
    ], [3.5, 8.6, 4.2])
    doc.add_paragraph('Git 服务为解析状态还会运行分支查询等命令；Terminal 会展示实际命令而不是只显示简写。')

    doc.add_heading('四、Terminal 输入与学习模式', 1)
    doc.add_paragraph('下方 Terminal 既是日志窗口，也是可以输入命令的工作区。选中项目后，终端默认进入该项目路径。')
    table(doc, ['操作', '效果'], [
        ('输入 git status + Enter', '执行并显示 Git 状态'),
        ('输入 cd .. + Enter', '切换 Terminal 会话目录，提示符更新'),
        ('输入 pwd + Enter', '输出当前目录'),
        ('↑ / ↓', '查看上一条/下一条历史命令'),
        ('历史', '弹出 SQLite 历史列表（含 UI / YOU 来源）'),
        ('停止', '停止最近启动的命令'),
        ('复制 / 清空', '复制选中或全部输出 / 清空输出显示'),
        ('Settings → Git 学习模式', '开关图形界面操作的中文命令解释'),
    ], [5.4, 10.7])
    doc.add_heading('4.1 自己输入命令', 2)
    code(doc, 'D:\\Projects\\Demo> git status\n'
              'On branch main\n'
              'nothing to commit, working tree clean\n'
              '  ↳ exit 0 · 0.04s')
    note(doc, '终端限制', '常见 Git/Python 等命令直接执行；CMD 内置命令与管道使用独立 CMD 执行器；cd 会保留目录，但普通 CMD 变量修改不会跨命令保留。此终端不是 ConPTY，不能运行 vim / python 交互 REPL 等依赖持续 TTY 的程序。')
    doc.add_paragraph('在 Terminal 里自己输入 `cd` 改变的是当前 Terminal 会话目录，不会暗中替换 Git 页的“当前项目”；若要对不同仓库执行 GUI Git 操作，先通过 Projects 切换项目。')

    doc.add_heading('五、其他模块操作', 1)
    table(doc, ['模块', '操作示例', '前提 / 注意'], [
        ('Git LFS', '检测 Git LFS、追踪 *.onnx、生成 .gitattributes', '本机需有 git-lfs；不会迁移历史'),
        ('Environment', '检测 Python/Conda/JDK/Node/Docker', '需对应工具安装于本机'),
        ('Export/Restore', '导出 requirements.txt、恢复依赖', '执行恢复前确认第三方依赖来源'),
        ('Docker', '查看容器/镜像，启动/停止，Compose Up', 'Docker Engine 可访问'),
        ('SSH', '检测公钥、生成 key、测试 GitHub/GitLab', '绝不复制私钥到第三方网站'),
        ('GitHub', 'Clone、Remote、创建 GitHub 模板/辅助仓库创建', '远程功能需网络/认证'),
        ('Generator', '创建 Python/Java/Node/AI 项目骨架', '不会覆盖已有非空目录'),
        ('Health / Ignore', '健康检查和合并 .gitignore', '输出建议需要自行审查'),
        ('Plugins', '安装/启用/禁用本地插件', '仅元数据管理，不自动运行插件代码'),
    ], [3.4, 7.3, 5.6])
    doc.add_heading('六、安全与常见问题', 1)
    table(doc, ['现象 / 疑问', '解释 / 解决'], [
        ('fatal: not a git repository', '先选择正确目录，或点击“初始化 Git”。'),
        ('git not found', '安装 Git for Windows，确保 PATH 能找到 git.exe，重启软件。'),
        ('Commit 显示 nothing to commit', '检查变更是否已 Stage；也可能文件没有修改。'),
        ('Push 提示 no upstream branch', '先设置 Remote 和跟踪分支，可在 Terminal 查看 git push -u origin main。'),
        ('远程认证失败', '配置 SSH/Windows Credential Manager/GitHub CLI；不在软件中保存密码。'),
        ('Docker 命令失败', '确认 Docker Desktop/Engine 已启动且当前账户有权限。'),
        ('Terminal 命令没有输出', '很多成功命令（如 git add）本身 stdout 为空，以 exit 0 判定成功。'),
        ('误操作风险', 'Discard、Stash Pop 等需要关注工作区状态；重要数据请先备份。'),
        ('升级失败', '查看 install.log，关闭旧程序再重试。'),
    ], [5.1, 11])
    note(doc, '安全提醒', '执行 Git、Docker、包管理器命令会改变实际项目；危险命令只做启发式提示，不是万能防护。输入任意命令前应确认内容可信。')
    doc.add_heading('七、发布与验收状态', 1)
    doc.add_paragraph('当前包是从需求和 UI 参考重建的 v1.1 源码与 Windows 原生在线引导安装器。项目包含 Python 服务测试、Git 临时仓库测试、SQLite 迁移测试及 Go 安装器模块测试。')
    doc.add_paragraph('构建所在环境为 Linux，无法下载 PySide6，因此不能声称已在 Windows 真机运行完整 GUI。首次使用建议先在测试机核验：安装 → 打开项目 → Stage → Commit → Terminal 手动 status → SQLite 历史 → 卸载。')
    doc.add_heading('八、参考资料', 1)
    for x in ['Git 命令手册：https://git-scm.com/docs',
              'Git add：https://git-scm.com/docs/git-add',
              'Git commit：https://git-scm.com/docs/git-commit',
              'Qt for Python：https://doc.qt.io/qtforpython-6/',
              'Qt QProcess：https://doc.qt.io/qt-6/qprocess.html']:
        doc.add_paragraph(x)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == '__main__':
    import sys
    generate(Path(sys.argv[1]))
